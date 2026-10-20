import io
import re
from pathlib import Path
from typing import Union, BinaryIO
import docx
from app.schemas.document import ParsedDocument, ParsedSection

def clean_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()

def parse_docx(
    source: Union[str, Path, BinaryIO, bytes],
    doc_id: str,
    filename: str = "document.docx",
) -> ParsedDocument:
    """
    Parses DOCX files into structural sections based on headings and tables.
    """
    if isinstance(source, bytes):
        doc = docx.Document(io.BytesIO(source))
    elif isinstance(source, (str, Path)):
        doc = docx.Document(str(source))
    else:
        doc = docx.Document(source)

    sections = []
    current_title = "Document Start"
    current_paragraphs = []
    full_text_parts = []

    def flush_section():
        nonlocal current_paragraphs, current_title
        if current_paragraphs:
            combined = "\n\n".join(current_paragraphs).strip()
            if combined:
                sections.append(
                    ParsedSection(
                        text=combined,
                        section_title=current_title,
                        metadata={"type": "text_block"}
                    )
                )
                full_text_parts.append(f"### {current_title}\n{combined}")
            current_paragraphs = []

    for para in doc.paragraphs:
        text = clean_text(para.text)
        if not text:
            continue

        style_name = para.style.name.lower() if para.style and para.style.name else ""
        if "heading" in style_name:
            flush_section()
            current_title = text
        else:
            current_paragraphs.append(text)

    flush_section()

    # Also extract structured tables if any
    for table_idx, table in enumerate(doc.tables, start=1):
        rows_data = []
        for row in table.rows:
            row_cells = [clean_text(cell.text) for cell in row.cells]
            if any(row_cells):
                rows_data.append(" | ".join(row_cells))
        if rows_data:
            table_text = "\n".join(rows_data)
            sections.append(
                ParsedSection(
                    text=table_text,
                    section_title=f"Table {table_idx}",
                    metadata={"type": "table", "rows": len(rows_data)}
                )
            )
            full_text_parts.append(f"### Table {table_idx}\n{table_text}")

    return ParsedDocument(
        doc_id=doc_id,
        filename=filename,
        source_type="docx",
        source_url=None,
        sections=sections,
        total_pages=None,
        raw_text="\n\n".join(full_text_parts),
        metadata={"sections_count": len(sections), "tables_count": len(doc.tables)}
    )
