import io
import sys
from pathlib import Path

# Fix Windows console UTF-8 output
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

# Add backend to sys.path
backend_root = Path(__file__).parent.parent
sys.path.insert(0, str(backend_root))

from app.services.parsers.txt_parser import parse_txt
from app.services.parsers.pdf_parser import parse_pdf
from app.services.parsers.docx_parser import parse_docx
from app.services.parsers.web_parser import parse_web_page
from app.services.parsers.dispatcher import parse_document, parse_url
from app.schemas.document import ParsedDocument

def test_txt_parser():
    sample_md = """# Architecture Overview
This is the introduction to DocuNexus.

## Ingestion Pipeline
The ingestion pipeline processes multiple document formats.

### PDF Support
Page numbers are tracked precisely.
"""
    doc = parse_txt(sample_md, doc_id="test_doc_1", filename="readme.md")
    assert isinstance(doc, ParsedDocument)
    assert doc.doc_id == "test_doc_1"
    assert doc.source_type == "txt"
    assert len(doc.sections) >= 3
    print("[PASS] TXT / Markdown parser test passed!")

def test_pdf_generation_and_parsing():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    doc = parse_pdf(pdf_bytes, doc_id="test_pdf_1", filename="sample.pdf")
    assert doc.total_pages == 2
    assert doc.source_type == "pdf"
    print("[PASS] PDF parser test passed!")

def test_docx_generation_and_parsing():
    import docx
    d = docx.Document()
    d.add_heading("Financial Summary", level=1)
    d.add_paragraph("Revenue grew by 24% year over year.")
    d.add_heading("Operational Highlights", level=2)
    d.add_paragraph("Customer retention remained at 98%.")
    
    buf = io.BytesIO()
    d.save(buf)
    docx_bytes = buf.getvalue()

    doc = parse_docx(docx_bytes, doc_id="test_docx_1", filename="report.docx")
    assert doc.source_type == "docx"
    assert len(doc.sections) == 2
    assert "Revenue grew by 24%" in doc.raw_text
    print("[PASS] DOCX parser test passed!")

def test_dispatcher_routing():
    # Test dispatcher routes by extension
    txt_doc = parse_document("Some raw note text", filename="meeting_notes.txt")
    assert txt_doc.source_type == "txt"
    print("[PASS] Dispatcher routing test passed!")

if __name__ == "__main__":
    print("Running parser unit tests...")
    test_txt_parser()
    test_pdf_generation_and_parsing()
    test_docx_generation_and_parsing()
    test_dispatcher_routing()
    print("ALL parser tests completed successfully!")
